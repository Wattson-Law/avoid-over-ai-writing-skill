#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
from pathlib import Path
from typing import Any

sys.dont_write_bytecode = True
VENDORED_PYTHON_DOCX = (
    Path(__file__).resolve().parent.parent / "vendor" / "python_docx_1_2_0"
)
if VENDORED_PYTHON_DOCX.is_dir():
    sys.path.insert(0, str(VENDORED_PYTHON_DOCX))

DOCX_IMPORT_ERROR: Exception | None = None
try:
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
except Exception as exc:
    DOCX_IMPORT_ERROR = exc
    Document = None
    WD_STYLE_TYPE = None
    WD_ALIGN_PARAGRAPH = None
    WD_TAB_ALIGNMENT = None
    WD_TAB_LEADER = None
    OxmlElement = None
    qn = None
    Inches = None
    Pt = None
    RGBColor = None


TOC_TITLE = "目录"
PLACEHOLDER = "[[TOC]]"
TOC_INSTRUCTION = ' TOC \\o "1-3" \\h \\z \\u '
HEADING_STYLE_RE = re.compile(r"^(?:heading|标题|標題)\s*([1-3])", re.IGNORECASE)
STATIC_TOC_RE = re.compile(r"\t\s*\d+\s*$|\.{3,}\s*\d+\s*$")


def _visible_text(element) -> str:
    return "".join(node.text or "" for node in element.xpath(".//w:t"))


def _has_toc_instruction(element) -> bool:
    for node in element.xpath(".//w:instrText"):
        if re.search(r"(?:^|\s)TOC(?:\s|$)", node.text or "", re.IGNORECASE):
            return True
    for node in element.xpath(".//w:fldSimple"):
        if re.search(
            r"(?:^|\s)TOC(?:\s|$)", node.get(qn("w:instr")) or "", re.IGNORECASE
        ):
            return True
    return False


def _heading_level(paragraph) -> int | None:
    if paragraph._p.xpath("./w:pPr/w:outlineLvl/@w:val") == ["9"]:
        return None
    style = paragraph.style
    while style is not None:
        match = HEADING_STYLE_RE.match((style.name or "").strip())
        if match:
            return int(match.group(1))
        ppr = style._element.pPr
        outline = None if ppr is None else ppr.find(qn("w:outlineLvl"))
        if outline is not None:
            value = outline.get(qn("w:val"))
            if value is not None and value.isdigit() and int(value) <= 2:
                return int(value) + 1
        style = style.base_style
    ppr = paragraph._p.pPr
    outline = None if ppr is None else ppr.find(qn("w:outlineLvl"))
    if outline is not None:
        value = outline.get(qn("w:val"))
        if value is not None and value.isdigit() and int(value) <= 2:
            return int(value) + 1
    return None


def _first_heading(doc):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() and _heading_level(paragraph) is not None:
            return paragraph
    return None


def _clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def _add_toc_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = TOC_INSTRUCTION
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "TOC_PENDING_UPDATE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, placeholder, end))


def _set_run_font(run, east_asia: str, size_pt: float, bold: bool) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:eastAsia"), east_asia)
    color = rpr.find(qn("w:color"))
    if color is not None:
        for name in ("w:themeColor", "w:themeTint", "w:themeShade"):
            color.attrib.pop(qn(name), None)


def _format_title(paragraph) -> None:
    _clear_paragraph(paragraph)
    paragraph.style = paragraph.part.document.styles["Normal"]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Inches(0)
    paragraph.paragraph_format.left_indent = Inches(0)
    paragraph.paragraph_format.right_indent = Inches(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(12)
    paragraph.paragraph_format.line_spacing = 1.2
    paragraph.paragraph_format.keep_with_next = True
    _set_run_font(paragraph.add_run(TOC_TITLE), "宋体", 18, True)


def _style_toc_entries(doc, text_width) -> None:
    for level, left_indent in ((1, 0.0), (2, 0.28), (3, 0.56)):
        style_id = f"TOC{level}"
        style = next(
            (candidate for candidate in doc.styles if candidate.style_id.upper() == style_id),
            None,
        )
        if style is None:
            style = next(
                (
                    candidate
                    for candidate in doc.styles
                    if (candidate.name or "").strip().lower() == f"toc {level}"
                ),
                None,
            )
        if style is None:
            style = doc.styles.add_style(f"toc {level}", WD_STYLE_TYPE.PARAGRAPH)
        style._element.set(qn("w:styleId"), style_id)
        style.name = f"toc {level}"
        style.base_style = doc.styles["Normal"]
        style.font.name = "Times New Roman"
        style.font.size = Pt(12)
        style.font.bold = False
        style.font.color.rgb = RGBColor(0, 0, 0)
        rpr = style._element.get_or_add_rPr()
        rfonts = rpr.get_or_add_rFonts()
        rfonts.set(qn("w:ascii"), "Times New Roman")
        rfonts.set(qn("w:hAnsi"), "Times New Roman")
        rfonts.set(qn("w:eastAsia"), "宋体")
        color = rpr.find(qn("w:color"))
        if color is not None:
            color.set(qn("w:val"), "000000")
            for attr in ("w:themeColor", "w:themeTint", "w:themeShade"):
                color.attrib.pop(qn(attr), None)
        fmt = style.paragraph_format
        fmt.left_indent = Inches(left_indent)
        fmt.right_indent = Inches(0)
        fmt.first_line_indent = Inches(0)
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(0)
        fmt.line_spacing = 1.5
        fmt.tab_stops.clear_all()
        fmt.tab_stops.add_tab_stop(
            text_width, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
        )


def _set_update_fields(doc) -> None:
    settings = doc.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.insert(0, update)
    update.set(qn("w:val"), "true")


def _paragraph_has_section_break(element) -> bool:
    return bool(element.xpath("./w:pPr/w:sectPr"))


def _section_break_starts_new_page(element) -> bool:
    sections = element.xpath("./w:pPr/w:sectPr")
    if not sections:
        return False
    section_type = sections[0].find(qn("w:type"))
    value = "nextPage" if section_type is None else section_type.get(qn("w:val"))
    return value not in {"continuous"}


def _paragraph_has_page_control(element) -> bool:
    return bool(
        element.xpath("./w:pPr/w:pageBreakBefore")
        or element.xpath(".//w:br[@w:type='page']")
        or _section_break_starts_new_page(element)
    )


def _normalize_break_before_body(toc_element, first_heading) -> int:
    removed = 0
    has_section_page_break = False
    current = toc_element.getnext()
    while current is not None and current is not first_heading._p:
        following = current.getnext()
        if current.tag == qn("w:p"):
            if _section_break_starts_new_page(current):
                has_section_page_break = True
            elif not _visible_text(current).strip() and _paragraph_has_page_control(current):
                current.getparent().remove(current)
                removed += 1
        current = following
    first_heading.paragraph_format.page_break_before = not has_section_page_break
    return removed


def _remove_recognized_static_entries(title_element, stop_element) -> int:
    removed = 0
    current = title_element.getnext()
    while current is not None and current is not stop_element:
        following = current.getnext()
        if current.tag == qn("w:p") and not _paragraph_has_section_break(current):
            text = _visible_text(current).strip()
            style_nodes = current.xpath("./w:pPr/w:pStyle")
            style_id = style_nodes[0].get(qn("w:val")) if style_nodes else ""
            if text and (
                re.fullmatch(r"TOC[1-3]", style_id or "", re.IGNORECASE)
                or STATIC_TOC_RE.search(text)
            ):
                current.getparent().remove(current)
                removed += 1
        current = following
    return removed


def _rewrite_existing_field_codes(doc) -> int:
    count = 0
    for node in doc.element.xpath(".//w:instrText"):
        if re.search(r"(?:^|\s)TOC(?:\s|$)", node.text or "", re.IGNORECASE):
            node.text = TOC_INSTRUCTION
            count += 1
    for node in doc.element.xpath(".//w:fldSimple"):
        instruction = node.get(qn("w:instr")) or ""
        if re.search(r"(?:^|\s)TOC(?:\s|$)", instruction, re.IGNORECASE):
            node.set(qn("w:instr"), TOC_INSTRUCTION.strip())
            count += 1
    return count


def ensure_toc(input_path: Path, output_path: Path) -> dict[str, Any]:
    if DOCX_IMPORT_ERROR is not None:
        raise RuntimeError(
            "vendored python-docx could not load; a Python runtime with lxml is required: "
            f"{type(DOCX_IMPORT_ERROR).__name__}: {DOCX_IMPORT_ERROR}"
        )
    if not input_path.is_file():
        raise FileNotFoundError(input_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if input_path.resolve() != output_path.resolve():
        shutil.copyfile(input_path, output_path)

    doc = Document(output_path)
    first_heading = _first_heading(doc)
    if first_heading is None:
        raise RuntimeError(
            "no Heading 1-3 or outline-level 0-2 paragraph was found; style headings before generating the TOC"
        )

    body = doc._element.body
    body_children = list(body.iterchildren())
    body_positions = {id(child): index for index, child in enumerate(body_children)}
    heading_position = body_positions[id(first_heading._p)]
    title_candidates = [
        paragraph
        for paragraph in doc.paragraphs
        if paragraph.text.strip() == TOC_TITLE
        and body_positions.get(id(paragraph._p), heading_position + 1) < heading_position
    ]
    placeholder = next(
        (paragraph for paragraph in doc.paragraphs if paragraph.text.strip() == PLACEHOLDER),
        None,
    )
    toc_children = [child for child in body_children if _has_toc_instruction(child)]

    if len(toc_children) > 1:
        raise RuntimeError(f"multiple TOC fields found ({len(toc_children)}); keep exactly one")

    created = False
    removed_static_entries = 0
    if toc_children:
        toc_element = toc_children[0]
        title = title_candidates[-1] if title_candidates else doc.add_paragraph()
        if not title_candidates:
            toc_element.addprevious(title._p)
    else:
        if placeholder is not None:
            field_paragraph = placeholder
            _clear_paragraph(field_paragraph)
        else:
            field_paragraph = doc.add_paragraph()
            first_heading._p.addprevious(field_paragraph._p)
        _add_toc_field(field_paragraph)
        toc_element = field_paragraph._p
        created = True
        title = title_candidates[-1] if title_candidates else doc.add_paragraph()
        if title_candidates:
            removed_static_entries = _remove_recognized_static_entries(
                title._p, first_heading._p
            )
        else:
            toc_element.addprevious(title._p)
        if title._p.getnext() is not toc_element:
            toc_element.getparent().remove(toc_element)
            title._p.addnext(toc_element)

    _format_title(title)
    preceding = title._p.getprevious()
    if preceding is None or not _paragraph_has_page_control(preceding):
        title.paragraph_format.page_break_before = True
    removed_breaks = _normalize_break_before_body(toc_element, first_heading)

    rewritten_fields = _rewrite_existing_field_codes(doc)
    section = doc.sections[-1]
    text_width = section.page_width - section.left_margin - section.right_margin
    _style_toc_entries(doc, text_width)
    _set_update_fields(doc)
    doc.save(output_path)

    return {
        "ok": True,
        "input": str(input_path.resolve()),
        "output": str(output_path.resolve()),
        "created": created,
        "toc_fields": max(1, rewritten_fields),
        "removed_static_entries": removed_static_entries,
        "removed_redundant_page_controls": removed_breaks,
        "needs_field_update": True,
    }


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Insert or normalize a three-level Word TOC before the first real heading."
    )
    parser.add_argument("docx", type=Path)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    try:
        report = ensure_toc(args.docx, args.out)
    except Exception as exc:
        report = {
            "ok": False,
            "input": str(args.docx),
            "output": str(args.out),
            "failures": [
                {"check": "toc_generation", "message": f"{type(exc).__name__}: {exc}"}
            ],
        }
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
