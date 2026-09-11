#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
import sys
import zipfile
from collections import Counter
from pathlib import Path

sys.dont_write_bytecode = True
VENDORED_PYTHON_DOCX = (
    Path(__file__).resolve().parent.parent / "vendor" / "python_docx_1_2_0"
)
if VENDORED_PYTHON_DOCX.is_dir():
    sys.path.insert(0, str(VENDORED_PYTHON_DOCX))

from docx import Document
from docx.oxml.ns import qn


OUTLINE_RE = re.compile(
    r"^(?:第[一二三四五六七八九十]+[章节部分]|"
    r"[一二三四五六七八九十]+[、.]|"
    r"\d+(?:\.\d+){0,3}[、.\s]|"
    r"[（(]?\d+[）)])"
)


def twips(length) -> int | None:
    return None if length is None else round(length.pt * 20)


def run_font(run) -> dict:
    rpr = run._element.rPr
    result = {
        "text": run.text,
        "size_pt": None if run.font.size is None else run.font.size.pt,
        "bold": run.bold,
        "italic": run.italic,
        "api_name": run.font.name,
    }
    if rpr is not None and rpr.rFonts is not None:
        for key in ("ascii", "hAnsi", "eastAsia", "cs"):
            result[key] = rpr.rFonts.get(qn(f"w:{key}"))
    return result


def paragraph_record(index: int, paragraph) -> dict:
    pf = paragraph.paragraph_format
    style = paragraph.style.name if paragraph.style else None
    ppr = paragraph._p.pPr
    return {
        "index": index,
        "text": paragraph.text,
        "style": style,
        "outline_candidate": bool(
            (style and style.lower().startswith("heading"))
            or OUTLINE_RE.match(paragraph.text.strip())
        ),
        "alignment": None if paragraph.alignment is None else str(paragraph.alignment),
        "first_line_indent_twips": twips(pf.first_line_indent),
        "left_indent_twips": twips(pf.left_indent),
        "right_indent_twips": twips(pf.right_indent),
        "space_before_pt": None if pf.space_before is None else pf.space_before.pt,
        "space_after_pt": None if pf.space_after is None else pf.space_after.pt,
        "line_spacing": pf.line_spacing,
        "section_break": bool(ppr is not None and ppr.find(qn("w:sectPr")) is not None),
        "runs": [run_font(run) for run in paragraph.runs if run.text],
    }


def table_record(index: int, table) -> dict:
    grid = table._tbl.tblGrid
    widths = [
        int(col.get(qn("w:w"), "0"))
        for col in grid.findall(qn("w:gridCol"))
    ]
    return {
        "index": index,
        "rows": len(table.rows),
        "columns": len(table.columns),
        "style": table.style.name if table.style else None,
        "grid_widths_twips": widths,
        "header": [cell.text for cell in table.rows[0].cells] if table.rows else [],
    }


def package_record(path: Path) -> dict:
    with zipfile.ZipFile(path) as package:
        names = package.namelist()
        document_xml = package.read("word/document.xml").decode("utf-8", "ignore")
    return {
        "media_files": sum(name.startswith("word/media/") for name in names),
        "has_comments": "word/comments.xml" in names,
        "tracked_insertions": len(re.findall(r"<w:ins(?:\s|>)", document_xml)),
        "tracked_deletions": len(re.findall(r"<w:del(?:\s|>)", document_xml)),
    }


def inspect(path: Path) -> dict:
    doc = Document(path)
    paragraphs = [
        paragraph_record(index, paragraph)
        for index, paragraph in enumerate(doc.paragraphs)
        if paragraph.text or (
            paragraph._p.pPr is not None
            and paragraph._p.pPr.find(qn("w:sectPr")) is not None
        )
    ]

    style_counts = Counter()
    font_counts = Counter()
    size_counts = Counter()
    for paragraph in doc.paragraphs:
        if paragraph.text:
            style_counts[paragraph.style.name if paragraph.style else "(none)"] += 1
        for run in paragraph.runs:
            if not run.text:
                continue
            rpr = run._element.rPr
            east_asia = None
            if rpr is not None and rpr.rFonts is not None:
                east_asia = rpr.rFonts.get(qn("w:eastAsia"))
            font_counts[east_asia or run.font.name or "(inherit)"] += len(run.text)
            size_counts[
                "(inherit)" if run.font.size is None else f"{run.font.size.pt:g}"
            ] += len(run.text)

    sections = []
    for index, section in enumerate(doc.sections):
        sections.append(
            {
                "index": index,
                "page_width_in": section.page_width.inches,
                "page_height_in": section.page_height.inches,
                "orientation": str(section.orientation),
                "start_type": str(section.start_type),
                "top_margin_in": section.top_margin.inches,
                "bottom_margin_in": section.bottom_margin.inches,
                "left_margin_in": section.left_margin.inches,
                "right_margin_in": section.right_margin.inches,
                "header_text": "\n".join(p.text for p in section.header.paragraphs),
                "footer_text": "\n".join(p.text for p in section.footer.paragraphs),
            }
        )

    return {
        "path": str(path.resolve()),
        "sections": sections,
        "paragraph_count": len(doc.paragraphs),
        "nonempty_paragraph_count": sum(bool(p.text.strip()) for p in doc.paragraphs),
        "paragraphs": paragraphs,
        "tables": [table_record(i, table) for i, table in enumerate(doc.tables)],
        "inline_shapes": len(doc.inline_shapes),
        "style_counts": style_counts.most_common(),
        "font_counts_by_character": font_counts.most_common(),
        "size_counts_by_character": size_counts.most_common(),
        "package": package_record(path),
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Inspect DOCX structure and formatting.")
    parser.add_argument("input", type=Path)
    parser.add_argument("--out", type=Path)
    args = parser.parse_args()

    report = inspect(args.input)
    payload = json.dumps(report, ensure_ascii=False, indent=2)
    if args.out:
        args.out.parent.mkdir(parents=True, exist_ok=True)
        args.out.write_text(payload + "\n", encoding="utf-8")
    else:
        print(payload)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
