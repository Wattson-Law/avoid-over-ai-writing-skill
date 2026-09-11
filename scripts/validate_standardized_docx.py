#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
import sys
import zipfile
from pathlib import Path
from typing import Any, Iterable
from xml.etree import ElementTree as ET

sys.dont_write_bytecode = True

VENDORED_PYTHON_DOCX = (
    Path(__file__).resolve().parent.parent / "vendor" / "python_docx_1_2_0"
)
if VENDORED_PYTHON_DOCX.is_dir():
    sys.path.insert(0, str(VENDORED_PYTHON_DOCX))

DOCX_IMPORT_ERROR: Exception | None = None
try:
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
except Exception as exc:  # defer optional dependency failures to structured JSON
    DOCX_IMPORT_ERROR = exc
    Document = None
    WD_CELL_VERTICAL_ALIGNMENT = None
    WD_ALIGN_PARAGRAPH = None
    qn = None


CAPTION_RE = re.compile(r"^表(\d+)$")
DATE_RE = re.compile(r"(?:19|20)\d{2}\s*年(?:\s*\d{1,2}\s*月)?")
TOP_HEADING_RE = re.compile(r"^(?:第[一二三四五六七八九十百]+[章节]|[一二三四五六七八九十百]+、)")
PAREN_HEADING_RE = re.compile(r"^[（(][一二三四五六七八九十百]+[）)]")
DECIMAL_HEADING_RE = re.compile(r"^(\d+(?:\.\d+){1,3})(?=\s|[^\d.])")
SINGLE_NUMBER_HEADING_RE = re.compile(r"^\d+[、.]\s*")
DETAIL_HEADING_RE = re.compile(r"^\d+[）)]\s*")
HEADING_STYLE_RE = re.compile(r"^(?:heading|标题|標題)\s*([1-9])", re.IGNORECASE)
TOC_STYLE_RE = re.compile(r"^(?:toc|目录|目錄)\s*\d+", re.IGNORECASE)
STRUCTURAL_PREFIX_RE = re.compile(
    r"^(?:第[一二三四五六七八九十百]+[章节]\s*|"
    r"[一二三四五六七八九十百]+、\s*|"
    r"[（(][一二三四五六七八九十百]+[）)]\s*|"
    r"\d+(?:\.\d+){0,3}[、.]\s*)"
)
AI_STYLE_RE = re.compile(
    r"前所未有|立足.{0,12}(?:设计|实际|全局)|坚持.{0,12}(?:思维|导向|原则)|"
    r"突出.{0,12}(?:价值|引领|重点)|致力于|持续释放|注入新动能|"
    r"构建.{0,20}新范式|多维度|全视角|坚实支撑|"
    r"从蓝图走向落地|从规划转化为成效|相互交织[、，]彼此影响|"
    r"赋能|全面提升|深度融合"
)
TOC_FIELD_RE = re.compile(r"(?:^|\s)TOC(?:\s|$)", re.IGNORECASE)
TOC_PLACEHOLDER_RE = re.compile(
    r"\[\[TOC\]\]|TOC_PENDING_UPDATE|TOC will populate|请(?:右键)?更新(?:目录|域)",
    re.IGNORECASE,
)

WHITE_FILLS = {None, "", "auto", "FFFFFF", "ffffff"}
TRUE_VALUES = {None, "", "1", "true", "on"}

R_AND_D_MODES = {"rd-application", "rd-implementation-outline"}
STRUCTURAL_TABLE_RE = re.compile(
    r"审核意见|审查意见|审批意见|签字|签名|签章|盖章|申报意见|协作单位意见"
)
FORM_FIELD_RE = re.compile(
    r"项目编号|项目名称|申报单位|承担单位|协作单位|依托平台|项目负责人|"
    r"负责人|起止年限|起止时间|申请经费|项目经费"
)
DATA_HEADER_RE = re.compile(
    r"序号|类别|识别项|成果形式|成果数量|单位名称|任务分工|研究工作内容|"
    r"考核目标|姓名|年龄|职务|职称|项目职务|项目分工|科目|数量|单价|"
    r"合计|参数细节|备注"
)

MODE_CHOICES = (
    "requirements-list",
    "decision-proposal",
    "rd-application",
    "rd-implementation-outline",
)
STRENGTH_CHOICES = ("format-only", "editorial", "restructure")

WORDPROCESSINGML_NAMESPACES = {
    "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "http://purl.oclc.org/ooxml/wordprocessingml/main",
}
COMMENT_ANCHOR_NAMES = {
    "commentRangeStart",
    "commentRangeEnd",
    "commentReference",
}
REVISION_NAMES = {
    "ins",
    "del",
    "moveFrom",
    "moveTo",
    "moveFromRangeStart",
    "moveFromRangeEnd",
    "moveToRangeStart",
    "moveToRangeEnd",
    "cellIns",
    "cellDel",
    "cellMerge",
    "numberingChange",
    "tblGridChange",
}
REVISION_RANGE_RE = re.compile(
    r"(?:ins|del|movefrom|moveto)(?:rangestart|rangeend)$", re.IGNORECASE
)


def normalize_text(text: str) -> str:
    return re.sub(r"\s+", "", text.replace("％", "%"))


def normalize_editorial_text(text: str) -> str:
    text = normalize_text(text)
    return re.sub(r"[，。；：、！？,.!?;:\"'“”‘’（）()《》〈〉【】\[\]]", "", text)


def element_text(element) -> str:
    return "".join(node.text or "" for node in element.xpath(".//w:t"))


def value_pt(value) -> float | None:
    if value is None:
        return None
    if hasattr(value, "pt"):
        return float(value.pt)
    if isinstance(value, (int, float)):
        return float(value)
    return None


def effective_paragraph_value(paragraph, name: str):
    direct = getattr(paragraph.paragraph_format, name)
    if direct is not None:
        return direct
    if name == "first_line_indent":
        # Word may deduplicate a paragraph's hanging indent into its numbering
        # definition on save. Resolve that actual formatting before paragraph styles.
        ppr = paragraph._p.pPr
        numpr = None if ppr is None else ppr.numPr
        if numpr is not None and numpr.numId is not None:
            root = paragraph.part.numbering_part.element
            nums = root.xpath(f'./w:num[@w:numId="{numpr.numId.val}"]/w:abstractNumId/@w:val')
            level = 0 if numpr.ilvl is None else numpr.ilvl.val
            if nums:
                indents = root.xpath(f'./w:abstractNum[@w:abstractNumId="{nums[0]}"]/w:lvl[@w:ilvl="{level}"]/w:pPr/w:ind')
                if indents:
                    from docx.shared import Twips
                    hanging = indents[0].get(qn("w:hanging"))
                    first = indents[0].get(qn("w:firstLine"))
                    if hanging is not None:
                        return Twips(-int(hanging))
                    if first is not None:
                        return Twips(int(first))
    style = paragraph.style
    seen = set()
    while style is not None and style.style_id not in seen:
        seen.add(style.style_id)
        value = getattr(style.paragraph_format, name)
        if value is not None:
            return value
        style = style.base_style
    return None


def effective_alignment(paragraph):
    if paragraph.alignment is not None:
        return paragraph.alignment
    if paragraph.style is not None:
        return paragraph.style.paragraph_format.alignment
    return None


def run_rfonts(run) -> dict[str, str | None]:
    result = {"ascii": None, "hAnsi": None, "eastAsia": None, "cs": None}
    rpr = run._element.rPr
    if rpr is not None and rpr.rFonts is not None:
        for key in result:
            result[key] = rpr.rFonts.get(qn(f"w:{key}"))
    if result["eastAsia"] is not None:
        return result
    for style in run_style_chain(run):
        style_rpr = style._element.rPr
        style_fonts = None if style_rpr is None else style_rpr.find(qn("w:rFonts"))
        if style_fonts is None:
            continue
        for key in result:
            if result[key] is None:
                result[key] = style_fonts.get(qn(f"w:{key}"))
        if result["eastAsia"] is not None:
            break
    return result


def run_style_chain(run):
    seen: set[str] = set()
    paragraph = getattr(run, "_parent", None)
    for initial in (run.style, getattr(paragraph, "style", None)):
        style = initial
        while style is not None:
            marker = style.style_id
            if marker in seen:
                break
            seen.add(marker)
            yield style
            style = style.base_style


def effective_run_size(run, paragraph) -> float | None:
    if run.font.size is not None:
        return run.font.size.pt
    for style in run_style_chain(run):
        if style.font.size is not None:
            return style.font.size.pt
    return None


def effective_run_bold(run, paragraph) -> bool:
    if run.bold is not None:
        return bool(run.bold)
    for style in run_style_chain(run):
        if style.font.bold is not None:
            return bool(style.font.bold)
    return False


def color_record_from_rpr(rpr, source: str) -> dict[str, str | None] | None:
    if rpr is None:
        return None
    color = rpr.find(qn("w:color"))
    if color is None:
        return None
    return {
        "value": color.get(qn("w:val")),
        "theme": color.get(qn("w:themeColor")),
        "tint": color.get(qn("w:themeTint")),
        "shade": color.get(qn("w:themeShade")),
        "source": source,
    }


def run_color_record(run) -> dict[str, str | None]:
    direct = color_record_from_rpr(run._element.rPr, "run")
    if direct is not None:
        return direct
    for style in run_style_chain(run):
        inherited = color_record_from_rpr(style._element.rPr, f"style:{style.style_id}")
        if inherited is not None:
            return inherited
    return {"value": None, "theme": None, "tint": None, "shade": None, "source": None}


def run_is_black(run) -> bool:
    color = run_color_record(run)
    value = (color["value"] or "").upper()
    return value == "000000" and all(
        color[key] in {None, ""} for key in ("theme", "tint", "shade")
    )


def font_matches(run, family: str) -> bool:
    fonts = run_rfonts(run)
    east_asia = (fonts["eastAsia"] or "").lower().replace(" ", "")
    api_name = (run.font.name or "").lower().replace(" ", "")
    candidate = east_asia or api_name
    if family == "song":
        return "宋体" in candidate or "simsun" in candidate
    if family == "fangsong":
        return "仿宋" in candidate or "fangsong" in candidate
    if family == "fangsong-gb":
        return "仿宋" in candidate or "fangsong" in candidate
    return True


def visible_runs(paragraph) -> list:
    candidates = list(paragraph.runs)
    for hyperlink in getattr(paragraph, "hyperlinks", ()):
        candidates.extend(hyperlink.runs)
    seen: set[int] = set()
    result = []
    for run in candidates:
        marker = id(run._element)
        if marker in seen or not run.text or not run.text.strip():
            continue
        seen.add(marker)
        result.append(run)
    return result


def approx(value: float | None, expected: float, tolerance: float = 0.12) -> bool:
    return value is not None and abs(value - expected) <= tolerance


def paragraph_line_multiple(paragraph) -> float | None:
    value = effective_paragraph_value(paragraph, "line_spacing")
    if isinstance(value, (int, float)):
        return float(value)
    return None


def paragraph_has_keep_next(paragraph) -> bool:
    return effective_paragraph_value(paragraph, "keep_with_next") is True


def paragraph_style_name(paragraph) -> str:
    return "" if paragraph.style is None else (paragraph.style.name or "").strip()


def is_toc_paragraph(paragraph) -> bool:
    if TOC_STYLE_RE.match(paragraph_style_name(paragraph)):
        return True
    instructions = "".join(
        node.text or "" for node in paragraph._element.xpath(".//w:instrText")
    )
    return bool(re.search(r"(?:^|\s)TOC(?:\s|$)", instructions, re.IGNORECASE))


def paragraph_numbering_level(paragraph) -> int | None:
    ppr = paragraph._element.pPr
    if ppr is None or ppr.numPr is None or ppr.numPr.ilvl is None:
        return None
    try:
        return int(ppr.numPr.ilvl.val)
    except (TypeError, ValueError):
        return None


def paragraph_outline_level(paragraph) -> int | None:
    ppr = paragraph._element.pPr
    if ppr is not None:
        outline = ppr.find(qn("w:outlineLvl"))
        if outline is not None:
            try:
                value = int(outline.get(qn("w:val")))
                if 0 <= value <= 8:
                    return value + 1
            except (TypeError, ValueError):
                pass
    style = paragraph.style
    visited: set[str] = set()
    while style is not None and style.style_id not in visited:
        visited.add(style.style_id)
        match = HEADING_STYLE_RE.match((style.name or "").strip())
        if match:
            return int(match.group(1))
        style_ppr = style.element.pPr
        outline = None if style_ppr is None else style_ppr.find(qn("w:outlineLvl"))
        if outline is not None:
            try:
                value = int(outline.get(qn("w:val")))
                if 0 <= value <= 8:
                    return value + 1
            except (TypeError, ValueError):
                pass
        style = style.base_style
    return None


def heading_level(paragraph, mode: str) -> int | None:
    if is_toc_paragraph(paragraph):
        return None
    if paragraph_style_name(paragraph) in {"AOW List", "AOW Quote", "AOW Code", "Title"}:
        return None
    # An explicit body outline level takes precedence over numbering-shaped text.
    # Markdown lists and code are body content even if they begin with "1.1".
    outline = paragraph._element.xpath("./w:pPr/w:outlineLvl/@w:val")
    if outline == ["9"]:
        return None
    stripped = paragraph.text.strip()
    outline_level = paragraph_outline_level(paragraph)
    if outline_level is not None:
        return outline_level
    if TOP_HEADING_RE.match(stripped):
        return 1
    decimal = DECIMAL_HEADING_RE.match(stripped)
    numbering_level = paragraph_numbering_level(paragraph)
    if mode == "requirements-list":
        if SINGLE_NUMBER_HEADING_RE.match(stripped):
            return 2
        if PAREN_HEADING_RE.match(stripped):
            return 3
        if DETAIL_HEADING_RE.match(stripped):
            return 4
        if numbering_level is not None and len(stripped) <= 80 and not re.search(r"[。；！？]$", stripped):
            return min(numbering_level + 2, 4)
    if mode == "decision-proposal":
        if decimal:
            return min(decimal.group(1).count(".") + 1, 3)
        if SINGLE_NUMBER_HEADING_RE.match(stripped):
            return 1
    if mode in R_AND_D_MODES:
        if PAREN_HEADING_RE.match(stripped):
            return 2
        if decimal:
            return min(decimal.group(1).count(".") + 1, 3)
        if SINGLE_NUMBER_HEADING_RE.match(stripped) or DETAIL_HEADING_RE.match(stripped):
            return 3
    return None


def paragraph_format_issues(paragraph, expected: dict[str, Any]) -> list[str]:
    issues: list[str] = []
    runs = visible_runs(paragraph)
    if not runs:
        return ["paragraph has no visible runs"]
    for run in runs:
        size = effective_run_size(run, paragraph)
        if not approx(size, expected["size"]):
            issues.append(f"run {run.text!r} uses {size!r} pt instead of {expected['size']} pt")
        if expected.get("bold") and not effective_run_bold(run, paragraph):
            issues.append(f"run {run.text!r} is not bold")
        expected_fonts = expected.get("fonts") or (expected["font"],)
        if not any(font_matches(run, family) for family in expected_fonts):
            issues.append(
                f"run {run.text!r} does not use an allowed East Asian font "
                f"({', '.join(expected_fonts)})"
            )
        if not expected.get("allow_colored") and not run_is_black(run):
            issues.append(f"run {run.text!r} is not black")
    indent = value_pt(effective_paragraph_value(paragraph, "first_line_indent"))
    if expected.get("no_indent") and indent not in {None, 0.0}:
        issues.append(f"first-line indent is {indent:g} pt")
    if "first_indent" in expected and not approx(indent, expected["first_indent"], 0.2):
        issues.append(f"first-line indent is {indent!r} pt instead of {expected['first_indent']} pt")
    line = paragraph_line_multiple(paragraph)
    if "line" in expected and not approx(line, expected["line"], 0.03):
        issues.append(f"line spacing is {line!r} instead of {expected['line']}")
    before = value_pt(effective_paragraph_value(paragraph, "space_before"))
    after = value_pt(effective_paragraph_value(paragraph, "space_after"))
    if "before" in expected and not approx(before or 0.0, expected["before"], 0.2):
        issues.append(f"space before is {before!r} pt instead of {expected['before']} pt")
    if "after" in expected and not approx(after or 0.0, expected["after"], 0.2):
        issues.append(f"space after is {after!r} pt instead of {expected['after']} pt")
    if expected.get("keep_next") and not paragraph_has_keep_next(paragraph):
        issues.append("keep-with-next is not enabled")
    alignment = effective_alignment(paragraph)
    if expected.get("alignment") == "center" and alignment != WD_ALIGN_PARAGRAPH.CENTER:
        issues.append("paragraph is not centered")
    if expected.get("alignment") == "left" and alignment not in {None, WD_ALIGN_PARAGRAPH.LEFT}:
        issues.append("paragraph is not left-aligned")
    return issues


def add_check(report: dict[str, Any], name: str, issues: Iterable[str]) -> None:
    details = list(issues)
    report["checks"][name] = {"ok": not details, "details": details}
    if details:
        report["failures"].extend({"check": name, "message": item} for item in details)


def xml_tag_parts(tag: str) -> tuple[str | None, str]:
    if tag.startswith("{") and "}" in tag:
        namespace, local_name = tag[1:].split("}", 1)
        return namespace, local_name
    return None, tag.rsplit(":", 1)[-1]


def is_revision_element(local_name: str) -> bool:
    return (
        local_name in REVISION_NAMES
        or local_name.endswith("PrChange")
        or bool(REVISION_RANGE_RE.search(local_name))
    )


def check_package(path: Path) -> list[str]:
    issues: list[str] = []
    revision_hits: dict[str, set[str]] = {}
    comment_anchor_hits: dict[str, set[str]] = {}
    with zipfile.ZipFile(path) as package:
        names = set(package.namelist())
        comment_parts = sorted(
            name
            for name in names
            if name.startswith("word/")
            and name.rsplit("/", 1)[-1].lower().startswith("comments")
            and name.lower().endswith(".xml")
        )
        for name in sorted(names):
            if not name.startswith("word/") or not name.lower().endswith(".xml"):
                continue
            try:
                root = ET.fromstring(package.read(name))
            except ET.ParseError as exc:
                issues.append(f"{name} is not well-formed XML: {exc}")
                continue
            for element in root.iter():
                namespace, local_name = xml_tag_parts(element.tag)
                if namespace not in WORDPROCESSINGML_NAMESPACES:
                    continue
                if local_name in COMMENT_ANCHOR_NAMES:
                    comment_anchor_hits.setdefault(name, set()).add(local_name)
                if is_revision_element(local_name):
                    revision_hits.setdefault(name, set()).add(local_name)
    if comment_parts:
        issues.append("comment parts are present: " + ", ".join(comment_parts))
    if comment_anchor_hits:
        details = "; ".join(
            f"{name} ({', '.join(sorted(tags))})"
            for name, tags in sorted(comment_anchor_hits.items())
        )
        issues.append("comment anchors are present: " + details)
    if revision_hits:
        details = "; ".join(
            f"{name} ({', '.join(sorted(tags))})"
            for name, tags in sorted(revision_hits.items())
        )
        issues.append("tracked changes are present: " + details)
    return issues


def check_sections(doc, mode: str) -> list[str]:
    issues: list[str] = []
    if mode == "requirements-list" and len(doc.sections) != 1:
        issues.append(f"requirements-list must use one section, found {len(doc.sections)}")
    for index, section in enumerate(doc.sections, 1):
        width = section.page_width.inches
        height = section.page_height.inches
        short_side, long_side = sorted((width, height))
        if abs(short_side - 8.27) > 0.06 or abs(long_side - 11.69) > 0.06:
            issues.append(f"section {index} is not A4 ({width:.2f} x {height:.2f} in)")
            continue
        landscape = width > height
        if mode in {"requirements-list", "decision-proposal"} and landscape:
            issues.append(f"section {index} is landscape but this mode requires portrait")
        if landscape and mode in {"rd-application", "rd-implementation-outline"}:
            expected = (1.25, 1.25, 1.0, 1.0)
        else:
            expected = (1.0, 1.0, 1.25, 1.25)
        actual = (
            section.top_margin.inches,
            section.bottom_margin.inches,
            section.left_margin.inches,
            section.right_margin.inches,
        )
        labels = ("top", "bottom", "left", "right")
        for label, got, want in zip(labels, actual, expected):
            if abs(got - want) > 0.05:
                issues.append(f"section {index} {label} margin is {got:.2f} in instead of {want:.2f} in")
    return issues


def body_child_with_descendant(body, descendant):
    current = descendant
    while current is not None and current.getparent() is not body:
        current = current.getparent()
    return current


def text_with_tabs(paragraph_element) -> str:
    parts: list[str] = []
    for node in paragraph_element.iter():
        namespace, local_name = xml_tag_parts(node.tag)
        if namespace not in WORDPROCESSINGML_NAMESPACES:
            continue
        if local_name == "t":
            parts.append(node.text or "")
        elif local_name == "tab":
            parts.append("\t")
    return "".join(parts).strip()


def toc_field_elements(doc) -> list:
    found: list = []
    seen: set[int] = set()
    body = doc._element.body
    for node in doc._element.xpath(".//w:instrText"):
        if not TOC_FIELD_RE.search(node.text or ""):
            continue
        child = body_child_with_descendant(body, node)
        if child is not None and id(child) not in seen:
            seen.add(id(child))
            found.append(child)
    for node in doc._element.xpath(".//w:fldSimple"):
        if not TOC_FIELD_RE.search(node.get(qn("w:instr")) or ""):
            continue
        child = body_child_with_descendant(body, node)
        if child is not None and id(child) not in seen:
            seen.add(id(child))
            found.append(child)
    return found


def toc_field_instruction(element) -> str:
    parts = [node.text or "" for node in element.xpath(".//w:instrText")]
    parts.extend(node.get(qn("w:instr")) or "" for node in element.xpath(".//w:fldSimple"))
    return re.sub(r"\s+", " ", " ".join(parts)).strip()


def toc_style_issues(doc, level: int) -> list[str]:
    issues: list[str] = []
    name = f"TOC {level}"
    style = next(
        (candidate for candidate in doc.styles if candidate.style_id.upper() == f"TOC{level}"),
        None,
    )
    if style is None:
        return [f"{name} paragraph style is missing"]
    style_chain = []
    current = style
    while current is not None:
        style_chain.append(current)
        current = current.base_style

    def inherited_font_attribute(attribute: str) -> str | None:
        for candidate in style_chain:
            rpr = candidate._element.rPr
            rfonts = None if rpr is None else rpr.find(qn("w:rFonts"))
            if rfonts is None:
                continue
            value = rfonts.get(qn(f"w:{attribute}"))
            if value is not None:
                return value
        return None

    east_asia = inherited_font_attribute("eastAsia")
    ascii_font = inherited_font_attribute("ascii")
    hansi_font = inherited_font_attribute("hAnsi")
    if east_asia != "宋体" or ascii_font != "Times New Roman" or hansi_font != "Times New Roman":
        issues.append(
            f"{name} must use Song + Times New Roman, found eastAsia={east_asia!r}, "
            f"ascii={ascii_font!r}, hAnsi={hansi_font!r}"
        )
    size = next(
        (value_pt(candidate.font.size) for candidate in style_chain if candidate.font.size is not None),
        None,
    )
    if not approx(size, 12.0):
        issues.append(f"{name} uses {size!r} pt instead of 12 pt")
    bold = next(
        (candidate.font.bold for candidate in style_chain if candidate.font.bold is not None),
        None,
    )
    if bold not in {False, None}:
        issues.append(f"{name} must not be bold")
    color_record = next(
        (
            record
            for candidate in style_chain
            if (record := color_record_from_rpr(candidate._element.rPr, candidate.style_id))
            is not None
        ),
        None,
    )
    color_value = None if color_record is None else (color_record["value"] or "").upper()
    color_theme = None if color_record is None else color_record["theme"]
    if color_value != "000000" or color_theme not in {None, ""}:
        issues.append(f"{name} must declare explicit black text")
    fmt = style.paragraph_format
    expected_indent = {1: 0.0, 2: 20.16, 3: 40.32}[level]
    indent = value_pt(fmt.left_indent)
    if not approx(indent or 0.0, expected_indent, 0.4):
        issues.append(
            f"{name} left indent is {indent!r} pt instead of about {expected_indent:g} pt"
        )
    first_indent = value_pt(fmt.first_line_indent)
    if first_indent not in {None, 0.0}:
        issues.append(f"{name} must not use a first-line or hanging indent")
    line = next(
        (
            candidate.paragraph_format.line_spacing
            for candidate in style_chain
            if candidate.paragraph_format.line_spacing is not None
        ),
        None,
    )
    if not isinstance(line, (int, float)) or not approx(float(line), 1.5, 0.03):
        issues.append(f"{name} line spacing is {line!r} instead of 1.5")
    before = next(
        (
            value_pt(candidate.paragraph_format.space_before)
            for candidate in style_chain
            if candidate.paragraph_format.space_before is not None
        ),
        None,
    )
    after = next(
        (
            value_pt(candidate.paragraph_format.space_after)
            for candidate in style_chain
            if candidate.paragraph_format.space_after is not None
        ),
        None,
    )
    if not approx(before or 0.0, 0.0, 0.2) or not approx(after or 0.0, 0.0, 0.2):
        issues.append(f"{name} must use 0 pt spacing before and after")
    tabs = style._element.xpath("./w:pPr/w:tabs/w:tab")
    has_right_dot = any(
        tab.get(qn("w:val")) == "right" and tab.get(qn("w:leader")) == "dot"
        for tab in tabs
    )
    if not has_right_dot:
        issues.append(f"{name} has no right-aligned dot-leader tab stop")
    return issues


def check_toc(doc, mode: str, require_toc: bool = False) -> list[str]:
    issues: list[str] = []
    fields = toc_field_elements(doc)
    titles = [paragraph for paragraph in doc.paragraphs if paragraph.text.strip() == "目录"]
    required = require_toc or (mode == "decision-proposal" and any(
        paragraph_outline_level(p) in {1, 2, 3} for p in doc.paragraphs
    ))
    present = bool(fields or titles)
    if mode == "requirements-list" and present and not require_toc:
        return ["requirements-list must not add a TOC unless the current user explicitly requests one"]
    if not required and not present:
        return issues
    if len(fields) != 1:
        issues.append(f"expected exactly one TOC field, found {len(fields)}")
    if len(titles) != 1:
        issues.append(f"expected exactly one TOC title paragraph, found {len(titles)}")

    body = doc._element.body
    children = list(body.iterchildren())
    positions = {id(child): index for index, child in enumerate(children)}
    heading_elements = [
        paragraph._p
        for paragraph in doc.paragraphs
        if paragraph.text.strip() and heading_level(paragraph, mode) in {1, 2, 3}
    ]
    positioned_headings = [positions[id(element)] for element in heading_elements if id(element) in positions]
    first_heading_position = min(positioned_headings) if positioned_headings else None
    if titles:
        title = titles[0]
        for item in paragraph_format_issues(
            title,
            {
                "font": "song",
                "size": 18.0,
                "bold": True,
                "alignment": "center",
                "no_indent": True,
            },
        ):
            issues.append(f"TOC title: {item}")
        title_position = positions.get(id(title._p))
    else:
        title_position = None
    field_position = positions.get(id(fields[0])) if fields else None
    if (
        title_position is not None
        and field_position is not None
        and first_heading_position is not None
        and not (title_position < field_position < first_heading_position)
    ):
        issues.append("TOC must appear after its title and before the first body heading")

    if fields:
        instruction = toc_field_instruction(fields[0])
        if not re.search(r'\\o\s*["“]1-3["”]', instruction, re.IGNORECASE):
            issues.append(f"TOC field does not cover heading levels 1-3: {instruction!r}")
        for switch in (r"\\h", r"\\z", r"\\u"):
            if not re.search(switch + r"(?:\s|$)", instruction, re.IGNORECASE):
                issues.append(f"TOC field is missing {switch} switch")

    body_text = element_text(doc._element.body)
    if TOC_PLACEHOLDER_RE.search(body_text):
        issues.append("TOC still contains placeholder text and has not been updated")

    toc_entries: list[tuple[int, str]] = []
    for paragraph_element in doc._element.xpath(".//w:p"):
        style_nodes = paragraph_element.xpath("./w:pPr/w:pStyle")
        if not style_nodes:
            continue
        style_id = style_nodes[0].get(qn("w:val")) or ""
        match = re.fullmatch(r"TOC([1-3])", style_id, re.IGNORECASE)
        if match:
            toc_entries.append((int(match.group(1)), text_with_tabs(paragraph_element)))
    if not toc_entries:
        issues.append("TOC field has no visible updated entries")
    else:
        for level, text in toc_entries:
            if not re.search(r"\t\s*\d+\s*$", text):
                issues.append(f"TOC {level} entry has no tab-separated page number: {text!r}")
        entry_titles = [re.sub(r"\t\s*\d+\s*$", "", text).strip() for _, text in toc_entries]
        for paragraph in doc.paragraphs:
            level = heading_level(paragraph, mode)
            heading_text = paragraph.text.strip()
            if level not in {1, 2, 3} or not heading_text or heading_text == "目录":
                continue
            if not any(normalize_text(heading_text) == normalize_text(item) for item in entry_titles):
                issues.append(f"heading is missing from the updated TOC: {heading_text!r}")

    for level in (1, 2, 3):
        issues.extend(toc_style_issues(doc, level))
    return issues


def on_off_element_active(element) -> bool:
    value = element.get(qn("w:val"))
    return value is None or value.strip().lower() in {"1", "true", "on"}


def paragraph_page_control_count(paragraph_element) -> int:
    manual_breaks = len(paragraph_element.xpath(".//w:br[@w:type='page']"))
    before_breaks = sum(
        on_off_element_active(node)
        for node in paragraph_element.xpath("./w:pPr/w:pageBreakBefore")
    )
    return manual_breaks + before_breaks


def paragraph_has_page_control(paragraph_element) -> bool:
    return paragraph_page_control_count(paragraph_element) > 0


def paragraph_has_section_break(paragraph_element) -> bool:
    return bool(paragraph_element.xpath("./w:pPr/w:sectPr"))


def paragraph_has_payload(paragraph_element) -> bool:
    if element_text(paragraph_element).strip():
        return True
    return bool(paragraph_element.xpath(".//w:drawing|.//w:pict|.//w:object"))


def check_redundant_page_controls(doc) -> list[str]:
    issues: list[str] = []
    children = list(doc._element.body.iterchildren())
    previous_control: tuple[int, str] | None = None
    for index, child in enumerate(children):
        if child.tag != qn("w:p"):
            previous_control = None
            continue
        page_count = paragraph_page_control_count(child)
        has_page = paragraph_has_page_control(child)
        has_section = paragraph_has_section_break(child)
        if page_count > 1:
            issues.append(f"body element {index} contains {page_count} active page controls")
        if has_page and has_section:
            issues.append(f"body element {index} combines a page control and a section break")
        kinds = []
        if has_page:
            kinds.append("page")
        if has_section:
            kinds.append("section")
        if kinds:
            current_kind = "+".join(kinds)
            if previous_control is not None:
                previous_index, previous_kind = previous_control
                issues.append(
                    f"body elements {previous_index} and {index} stack consecutive "
                    f"{previous_kind}/{current_kind} controls across only blank paragraphs"
                )
            previous_control = (index, current_kind)
            if has_page and not has_section and paragraph_has_payload(child):
                previous_control = None
        elif paragraph_has_payload(child):
            previous_control = None
    return issues


def first_heading_index(doc, mode: str) -> int | None:
    for index, paragraph in enumerate(doc.paragraphs):
        if heading_level(paragraph, mode) is not None:
            return index
    return None


def is_explicit_body(paragraph) -> bool:
    """Markdown body blocks explicitly opt out of cover/title inference."""
    return (
        paragraph_style_name(paragraph) != "Title"
        and paragraph._p.xpath('./w:pPr/w:outlineLvl/@w:val') == ["9"]
    )


def check_cover_and_title(doc, mode: str, allow_colored_text: bool) -> list[str]:
    issues: list[str] = []
    nonempty = [(i, p) for i, p in enumerate(doc.paragraphs) if p.text.strip()]
    if not nonempty:
        return ["document has no visible paragraphs"]
    heading_index = first_heading_index(doc, mode)
    cover = [
        (i, p) for i, p in nonempty
        if (heading_index is None or i < heading_index)
        and not is_explicit_body(p) and not is_toc_paragraph(p)
    ]
    if mode in {"requirements-list", "decision-proposal"} and cover:
        title_index, title = cover[0]
        title_expected = {
            "font": "song", "size": 18.0, "bold": True, "alignment": "center",
            "allow_colored": allow_colored_text,
        }
        for item in paragraph_format_issues(title, title_expected):
            issues.append(f"title paragraph {title_index}: {item}")
    if mode == "decision-proposal" and cover:
        date_candidates = [(i, p) for i, p in cover if DATE_RE.search(p.text)]
        if date_candidates:
            date_index, date = date_candidates[-1]
            unit_candidates = [(i, p) for i, p in cover if i < date_index and i != cover[0][0]]
            expected = {
                "font": "song", "size": 16.0, "bold": True, "alignment": "center",
                "allow_colored": allow_colored_text,
            }
            for item in paragraph_format_issues(date, expected):
                issues.append(f"cover date paragraph {date_index}: {item}")
            if unit_candidates:
                unit_index, unit = unit_candidates[-1]
                for item in paragraph_format_issues(unit, expected):
                    issues.append(f"cover unit paragraph {unit_index}: {item}")
            else:
                issues.append("cover unit paragraph was not found before the date")
    return issues


def check_headings(
    doc, mode: str, allow_colored_text: bool, has_structural_tables: bool = False,
    require_headings: bool = True,
) -> list[str]:
    issues: list[str] = []
    found = 0
    for index, paragraph in enumerate(doc.paragraphs):
        level = heading_level(paragraph, mode)
        if level is None:
            continue
        found += 1
        if mode == "requirements-list" and level > 1:
            expected = {
                "font": "fangsong", "size": 16.0, "alignment": "left",
                "no_indent": True, "line": 1.5, "before": 0.0, "after": 0.0,
                "keep_next": True, "allow_colored": allow_colored_text,
            }
        elif mode in {"requirements-list", "decision-proposal"}:
            expected = {
                "font": "song", "size": 16.0, "bold": True, "alignment": "left",
                "no_indent": True, "line": 1.2, "before": 6.0, "after": 4.0,
                "keep_next": True, "allow_colored": allow_colored_text,
            }
        elif level == 1:
            expected = {
                "font": "fangsong-gb", "size": 16.0, "bold": True,
                "alignment": "left", "no_indent": True, "before": 12.0,
                "after": 12.0, "keep_next": True, "allow_colored": allow_colored_text,
            }
        else:
            expected = {
                "font": "fangsong", "size": 14.0, "bold": True,
                "alignment": "left", "no_indent": True, "keep_next": True,
                "allow_colored": allow_colored_text,
            }
        for item in paragraph_format_issues(paragraph, expected):
            issues.append(f"heading paragraph {index} ({paragraph.text!r}): {item}")
    if found == 0 and require_headings and not (mode in R_AND_D_MODES and has_structural_tables):
        issues.append("no numbered heading paragraphs were found")
    return issues


def check_body_paragraphs(doc, mode: str, allow_colored_text: bool) -> list[str]:
    issues: list[str] = []
    heading_index = first_heading_index(doc, mode)
    for index, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if not text or CAPTION_RE.fullmatch(text) or is_toc_paragraph(paragraph) or heading_level(paragraph, mode) is not None:
            continue
        if paragraph_style_name(paragraph) == "Title":
            continue
        if not is_explicit_body(paragraph) and (heading_index is None or index < heading_index):
            continue
        if mode in {"requirements-list", "decision-proposal"}:
            expected = {"font": "fangsong", "size": 16.0, "first_indent": 32.0, "line": 1.5, "allow_colored": allow_colored_text}
        else:
            expected = {"font": "fangsong", "size": 14.0, "first_indent": 28.0, "line": 1.5, "allow_colored": allow_colored_text}
        style_name = paragraph_style_name(paragraph)
        if style_name in {"AOW List", "AOW Quote"}:
            expected["first_indent"] = -12.0 if style_name == "AOW List" else 0.0
        elif style_name == "AOW Code":
            expected.update(font="song", size=12.0, first_indent=0.0, line=1.0)
        for item in paragraph_format_issues(paragraph, expected):
            issues.append(f"body paragraph {index}: {item}")
    return issues


def all_visible_paragraphs(doc) -> Iterable[tuple[str, Any]]:
    for index, paragraph in enumerate(doc.paragraphs):
        yield f"paragraph {index}", paragraph
    for table_index, table in enumerate(doc.tables):
        seen: set[int] = set()
        for row_index, row in enumerate(table.rows):
            for cell_index, cell in enumerate(row.cells):
                if id(cell._tc) in seen:
                    continue
                seen.add(id(cell._tc))
                for paragraph_index, paragraph in enumerate(cell.paragraphs):
                    yield f"table {table_index + 1} row {row_index + 1} cell {cell_index + 1} paragraph {paragraph_index + 1}", paragraph
    for section_index, section in enumerate(doc.sections):
        for area_name, area in (("header", section.header), ("footer", section.footer)):
            for paragraph_index, paragraph in enumerate(area.paragraphs):
                yield f"section {section_index + 1} {area_name} paragraph {paragraph_index + 1}", paragraph


def check_global_text_color(doc, allow_colored_text: bool) -> list[str]:
    if allow_colored_text:
        return []
    issues: list[str] = []
    for location, paragraph in all_visible_paragraphs(doc):
        for run in visible_runs(paragraph):
            if not run_is_black(run):
                color = run_color_record(run)
                issues.append(f"{location} run {run.text!r} uses non-black color {color}")
    return issues


def border_valid(border) -> bool:
    if border is None:
        return False
    value = border.get(qn("w:val"))
    size = border.get(qn("w:sz"))
    color = (border.get(qn("w:color")) or "").upper()
    has_theme_color = any(
        border.get(qn(name)) not in {None, ""}
        for name in ("w:themeColor", "w:themeTint", "w:themeShade")
    )
    return value == "single" and size == "4" and color == "000000" and not has_theme_color


def border_disabled(border) -> bool:
    return border is not None and (border.get(qn("w:val")) or "").lower() in {"nil", "none"}


def table_border_issues(table, table_number: int) -> list[str]:
    issues: list[str] = []
    tbl_borders = table._tbl.tblPr.find(qn("w:tblBorders"))
    required_table_edges = ("top", "left", "bottom", "right", "insideH", "insideV")
    if tbl_borders is not None:
        for edge in required_table_edges:
            if not border_valid(tbl_borders.find(qn(f"w:{edge}"))):
                issues.append(f"table {table_number} {edge} border is not a 0.5 pt black single line")
    seen: set[int] = set()
    for row_index, row in enumerate(table.rows):
        for cell_index, cell in enumerate(row.cells):
            if id(cell._tc) in seen:
                continue
            seen.add(id(cell._tc))
            borders = cell._tc.tcPr.find(qn("w:tcBorders"))
            if borders is None:
                if tbl_borders is None:
                    issues.append(f"table {table_number} row {row_index + 1} cell {cell_index + 1} has no explicit borders")
                continue
            for edge in ("top", "left", "bottom", "right"):
                node = borders.find(qn(f"w:{edge}"))
                if node is None and edge == "left":
                    node = borders.find(qn("w:start"))
                if node is None and edge == "right":
                    node = borders.find(qn("w:end"))
                if tbl_borders is not None and node is None:
                    continue
                if border_disabled(node):
                    issues.append(
                        f"table {table_number} row {row_index + 1} cell {cell_index + 1} "
                        f"{edge} border is explicitly disabled"
                    )
                    continue
                if not border_valid(node):
                    issues.append(
                        f"table {table_number} row {row_index + 1} cell {cell_index + 1} "
                        f"{edge} border is not a 0.5 pt black single line"
                    )
    return issues


def unique_table_cells(table):
    seen: set[int] = set()
    for row in table.rows:
        for cell in row.cells:
            marker = id(cell._tc)
            if marker in seen:
                continue
            seen.add(marker)
            yield cell


def table_text(table) -> str:
    return "\n".join(cell.text.strip() for cell in unique_table_cells(table) if cell.text.strip())


def table_first_row_text(table) -> str:
    if not table.rows:
        return ""
    seen: set[int] = set()
    parts: list[str] = []
    for cell in table.rows[0].cells:
        marker = id(cell._tc)
        if marker in seen:
            continue
        seen.add(marker)
        if cell.text.strip():
            parts.append(cell.text.strip())
    return " ".join(parts)


def structural_table_ids(doc, mode: str) -> set[int]:
    if mode not in R_AND_D_MODES:
        return set()
    children = list(doc._element.body.iterchildren())
    body_positions = {id(child): index for index, child in enumerate(children)}
    heading_positions = [
        body_positions[id(paragraph._element)]
        for paragraph in doc.paragraphs
        if id(paragraph._element) in body_positions and heading_level(paragraph, mode) is not None
    ]
    first_heading_position = min(heading_positions) if heading_positions else None
    structural: set[int] = set()
    for table in doc.tables:
        marker = id(table._tbl)
        position = body_positions.get(marker)
        text = table_text(table)
        first_row = table_first_row_text(table)
        before_body = (
            position is not None
            and first_heading_position is not None
            and position < first_heading_position
        )
        form_fields = len(set(FORM_FIELD_RE.findall(text)))
        looks_like_form = form_fields >= 2 and not DATA_HEADER_RE.search(first_row)
        looks_like_layout = len(table.rows) <= 1 and len(table.columns) <= 2 and not DATA_HEADER_RE.search(first_row)
        if before_body or STRUCTURAL_TABLE_RE.search(text) or looks_like_form or looks_like_layout:
            structural.add(marker)
    return structural


def cell_fill(cell) -> str | None:
    shading = cell._tc.tcPr.find(qn("w:shd"))
    return None if shading is None else shading.get(qn("w:fill"))


def cell_has_explicit_white_fill(cell) -> bool:
    shading = cell._tc.tcPr.find(qn("w:shd"))
    if shading is None:
        return False
    fill = (shading.get(qn("w:fill")) or "").upper()
    has_theme_fill = any(
        shading.get(qn(name)) not in {None, ""}
        for name in ("w:themeFill", "w:themeFillTint", "w:themeFillShade")
    )
    return fill == "FFFFFF" and not has_theme_fill


def check_table_captions(doc, structural_tables: set[int]) -> list[str]:
    issues: list[str] = []
    children = list(doc._element.body.iterchildren())
    expected = 1
    table_count = 0
    semantic_table_count = 0
    caption_texts = [
        paragraph.text.strip()
        for paragraph in doc.paragraphs
        if CAPTION_RE.fullmatch(paragraph.text.strip())
    ]
    for index, child in enumerate(children):
        if child.tag != qn("w:tbl"):
            continue
        table_count += 1
        if id(child) in structural_tables:
            continue
        semantic_table_count += 1
        if index == 0 or children[index - 1].tag != qn("w:p"):
            issues.append(f"table {table_count} is not immediately preceded by a caption paragraph")
            expected += 1
            continue
        text = element_text(children[index - 1]).strip()
        wanted = f"表{expected}"
        if text != wanted:
            issues.append(f"table {table_count} caption is {text!r}, expected {wanted!r}")
        expected += 1
    wanted_captions = [f"表{i}" for i in range(1, semantic_table_count + 1)]
    if caption_texts != wanted_captions:
        issues.append(f"caption sequence is {caption_texts!r}, expected {wanted_captions!r}")
    return issues


def check_caption_format(doc, mode: str) -> list[str]:
    issues: list[str] = []
    for index, paragraph in enumerate(doc.paragraphs):
        if not CAPTION_RE.fullmatch(paragraph.text.strip()):
            continue
        expected = {
            "font": "song",
            "fonts": ("song", "fangsong") if mode in R_AND_D_MODES else ("song",),
            "size": 12.0,
            "alignment": "center",
            "no_indent": True,
        }
        for item in paragraph_format_issues(paragraph, expected):
            issues.append(f"caption paragraph {index}: {item}")
    return issues


def check_tables(doc, mode: str, structural_tables: set[int]) -> list[str]:
    issues: list[str] = []
    max_text_width_twips = max(
        round((section.page_width.inches - section.left_margin.inches - section.right_margin.inches) * 1440)
        for section in doc.sections
    )
    for table_index, table in enumerate(doc.tables, 1):
        is_structural = id(table._tbl) in structural_tables
        issues.extend(table_border_issues(table, table_index))
        grid_width = sum(int(col.get(qn("w:w"), "0")) for col in table._tbl.tblGrid.findall(qn("w:gridCol")))
        if grid_width and grid_width > max_text_width_twips + 20:
            issues.append(f"table {table_index} width {grid_width} twips exceeds available text width")
        seen: set[int] = set()
        for row_index, row in enumerate(table.rows):
            for cell_index, cell in enumerate(row.cells):
                if id(cell._tc) in seen:
                    continue
                seen.add(id(cell._tc))
                if (
                    mode == "decision-proposal"
                    and cell.vertical_alignment != WD_CELL_VERTICAL_ALIGNMENT.CENTER
                ):
                    issues.append(
                        f"table {table_index} row {row_index + 1} cell {cell_index + 1} "
                        "is not vertically centered"
                    )
                fill = cell_fill(cell)
                if not is_structural and not cell_has_explicit_white_fill(cell):
                    issues.append(
                        f"table {table_index} row {row_index + 1} cell {cell_index + 1} "
                        f"must declare an explicit FFFFFF cell fill; found {fill!r}"
                    )
                for paragraph in cell.paragraphs:
                    if mode == "decision-proposal":
                        alignment = effective_alignment(paragraph)
                        if row_index == 0 and alignment != WD_ALIGN_PARAGRAPH.CENTER:
                            issues.append(
                                f"table {table_index} header cell {cell_index + 1} "
                                "paragraph is not horizontally centered"
                            )
                        elif row_index > 0 and visible_runs(paragraph) and alignment not in {
                            None,
                            WD_ALIGN_PARAGRAPH.LEFT,
                            WD_ALIGN_PARAGRAPH.CENTER,
                        }:
                            issues.append(
                                f"table {table_index} row {row_index + 1} cell {cell_index + 1} "
                                "paragraph must be left-aligned or centered"
                            )
                    for run in visible_runs(paragraph):
                        if not run_is_black(run):
                            issues.append(
                                f"table {table_index} row {row_index + 1} cell {cell_index + 1} "
                                f"run {run.text!r} is not black"
                            )
                        size = effective_run_size(run, paragraph)
                        if is_structural:
                            continue
                        if mode in {"requirements-list", "decision-proposal"}:
                            if not approx(size, 12.0):
                                issues.append(
                                    f"table {table_index} row {row_index + 1} cell {cell_index + 1} "
                                    f"run {run.text!r} uses {size!r} pt instead of 12 pt"
                                )
                            if not font_matches(run, "song"):
                                issues.append(
                                    f"table {table_index} row {row_index + 1} cell {cell_index + 1} "
                                    f"run {run.text!r} does not use Song"
                                )
                        elif size is None or size < 10.5 - 0.12 or size > 14.0 + 0.12:
                            issues.append(
                                f"table {table_index} row {row_index + 1} cell {cell_index + 1} "
                                f"run {run.text!r} uses {size!r} pt outside 10.5-14 pt"
                            )
        if table.rows and not is_structural:
            header = table.rows[0]
            tr_pr = header._tr.trPr
            header_marker = None if tr_pr is None else tr_pr.find(qn("w:tblHeader"))
            if header_marker is None or not on_off_element_active(header_marker):
                issues.append(f"table {table_index} header row is not marked to repeat")
            for cell_index, cell in enumerate(header.cells):
                for paragraph in cell.paragraphs:
                    for run in visible_runs(paragraph):
                        if not effective_run_bold(run, paragraph):
                            issues.append(f"table {table_index} header cell {cell_index + 1} run {run.text!r} is not bold")
    return issues


def narrative_segments(doc) -> list[str]:
    parts: list[str] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text or CAPTION_RE.fullmatch(text) or is_toc_paragraph(paragraph):
            continue
        text = STRUCTURAL_PREFIX_RE.sub("", text, count=1)
        if text.strip():
            parts.append(text.strip())
    for table in doc.tables:
        for cell in unique_table_cells(table):
            for paragraph in cell.paragraphs:
                text = paragraph.text.strip()
                if not text or CAPTION_RE.fullmatch(text):
                    continue
                text = STRUCTURAL_PREFIX_RE.sub("", text, count=1)
                if text.strip():
                    parts.append(text.strip())
    return parts


def narrative_text(doc) -> str:
    return "\n".join(normalize_text(part) for part in narrative_segments(doc))


def check_editorial_pass(source_path: Path | None, output_doc, strength: str) -> tuple[list[str], list[str]]:
    issues: list[str] = []
    warnings: list[str] = []
    if strength == "format-only":
        return issues, warnings
    if source_path is None:
        return [f"--source is required for {strength} validation"], warnings
    source_doc = Document(source_path)
    source_segments = narrative_segments(source_doc)
    output_segments = narrative_segments(output_doc)
    source_text = "\n".join(normalize_text(part) for part in source_segments)
    output_text = "\n".join(normalize_text(part) for part in output_segments)
    output_compact = "".join(normalize_editorial_text(part) for part in output_segments)
    unchanged_trigger_segments: list[str] = []
    for segment_index, segment in enumerate(source_segments, 1):
        normalized = normalize_editorial_text(segment)
        triggers = sorted(set(AI_STYLE_RE.findall(segment)))
        if not triggers or len(normalized) < 12 or normalized not in output_compact:
            continue
        preview = re.sub(r"\s+", " ", segment).strip()
        if len(preview) > 80:
            preview = preview[:77] + "..."
        unchanged_trigger_segments.append(
            f"source narrative segment {segment_index} remains verbatim "
            f"with triggers {', '.join(triggers[:6])}: {preview!r}"
        )
    if unchanged_trigger_segments:
        issues.extend(unchanged_trigger_segments[:25])
        if len(unchanged_trigger_segments) > 25:
            issues.append(
                f"{len(unchanged_trigger_segments) - 25} additional AI-style source segments remain verbatim"
            )
    if normalize_editorial_text(source_text) == normalize_editorial_text(output_text):
        triggers = sorted(set(AI_STYLE_RE.findall(source_text)))
        if triggers and not unchanged_trigger_segments:
            issues.append(
                "narrative text is unchanged despite clear editorial trigger phrases: "
                + ", ".join(triggers[:12])
            )
        elif not triggers:
            warnings.append("narrative text is unchanged; no configured editorial trigger phrase was detected")
    return issues, warnings


def validate(args) -> dict[str, Any]:
    if DOCX_IMPORT_ERROR is not None:
        raise RuntimeError(
            "vendored python-docx 1.2.0 could not load; install its runtime "
            f"dependencies (lxml and typing_extensions): {type(DOCX_IMPORT_ERROR).__name__}: "
            f"{DOCX_IMPORT_ERROR}"
        )
    doc = Document(args.output)
    structural_tables = structural_table_ids(doc, args.mode)
    report: dict[str, Any] = {
        "ok": False,
        "output": str(args.output.resolve()),
        "source": None if args.source is None else str(args.source.resolve()),
        "mode": args.mode,
        "processing_strength": args.processing_strength,
        "checks": {},
        "failures": [],
        "warnings": [],
        "note": "Internal format validation only; do not publish failures as a user-facing issue list.",
    }
    add_check(report, "clean_package", check_package(args.output))
    add_check(report, "page_layout", check_sections(doc, args.mode))
    add_check(report, "page_controls", check_redundant_page_controls(doc))
    add_check(report, "toc", check_toc(doc, args.mode, args.require_toc))
    add_check(report, "cover_and_title", check_cover_and_title(doc, args.mode, args.allow_colored_text))
    add_check(
        report,
        "heading_format",
        check_headings(doc, args.mode, args.allow_colored_text, bool(structural_tables),
                       require_headings=args.source is None or any(
                           heading_level(p, args.mode) is not None for p in Document(args.source).paragraphs
                       )),
    )
    add_check(report, "body_format", check_body_paragraphs(doc, args.mode, args.allow_colored_text))
    add_check(report, "global_text_color", check_global_text_color(doc, args.allow_colored_text))
    add_check(report, "table_captions", check_table_captions(doc, structural_tables))
    add_check(report, "caption_format", check_caption_format(doc, args.mode))
    add_check(report, "table_format", check_tables(doc, args.mode, structural_tables))
    editorial_issues, editorial_warnings = check_editorial_pass(
        args.source, doc, args.processing_strength
    )
    add_check(report, "editorial_pass", editorial_issues)
    report["warnings"].extend(editorial_warnings)
    report["ok"] = not report["failures"]
    return report


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Validate a DOCX standardized by the AOW technical-document skill."
    )
    parser.add_argument("output", type=Path)
    parser.add_argument("--mode", required=True, choices=MODE_CHOICES)
    parser.add_argument("--source", type=Path)
    parser.add_argument(
        "--processing-strength",
        choices=STRENGTH_CHOICES,
        default="editorial",
    )
    parser.add_argument("--allow-colored-text", action="store_true")
    parser.add_argument("--require-toc", action="store_true", help="Use only when the user explicitly requests a TOC outside mode defaults")
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()

    try:
        report = validate(args)
    except Exception as exc:  # keep failures machine-readable for host agents
        report = {
            "ok": False,
            "output": str(args.output),
            "source": None if args.source is None else str(args.source),
            "mode": args.mode,
            "processing_strength": args.processing_strength,
            "checks": {},
            "failures": [{"check": "validator_runtime", "message": f"{type(exc).__name__}: {exc}"}],
            "warnings": [],
            "note": "Internal format validation only; do not publish failures as a user-facing issue list.",
        }

    if args.report:
        try:
            payload = json.dumps(report, ensure_ascii=False, indent=2)
            args.report.parent.mkdir(parents=True, exist_ok=True)
            args.report.write_text(payload + "\n", encoding="utf-8")
        except Exception as exc:
            message = f"{type(exc).__name__}: {exc}"
            report.setdefault("checks", {})["report_write"] = {
                "ok": False,
                "details": [message],
            }
            report.setdefault("failures", []).append(
                {"check": "report_write", "message": message}
            )
            report["ok"] = False
    payload = json.dumps(report, ensure_ascii=False, indent=2)
    print(payload)
    if report["ok"]:
        print("Standardized DOCX validation passed.", file=sys.stderr)
        return 0
    print("Standardized DOCX validation failed.", file=sys.stderr)
    return 1


if __name__ == "__main__":
    raise SystemExit(main())
